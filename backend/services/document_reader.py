"""
EduGenAI - Document Reader Service
Extracts text from PowerPoint (PPTX) and Word (DOCX) files.
"""
import logging
import io
from docx import Document
from pptx import Presentation

logger = logging.getLogger(__name__)

async def extract_text_from_docx(file_bytes: bytes) -> dict:
    """Extracts text from a DOCX file."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)
        
        logger.info(f"DOCX extraction complete: {len(paragraphs)} paragraphs, {len(text)} chars.")
        
        return {
            "text": text,
            "method": "python-docx",
            "word_count": len(text.split()),
            "paragraph_count": len(paragraphs)
        }
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        raise RuntimeError(f"DOCX extraction failed: {str(e)}")


async def extract_text_from_pptx(file_bytes: bytes) -> dict:
    """Extracts text from a PPTX file."""
    try:
        prs = Presentation(io.BytesIO(file_bytes))
        slides_text = []
        
        for i, slide in enumerate(prs.slides):
            slide_content = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text.strip())
            
            if slide_content:
                slides_text.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_content))
                
        text = "\n\n".join(slides_text)
        
        logger.info(f"PPTX extraction complete: {len(prs.slides)} slides, {len(text)} chars.")
        
        return {
            "text": text,
            "method": "python-pptx",
            "word_count": len(text.split()),
            "slide_count": len(prs.slides)
        }
    except Exception as e:
        logger.error(f"PPTX extraction failed: {e}")
        raise RuntimeError(f"PPTX extraction failed: {str(e)}")
